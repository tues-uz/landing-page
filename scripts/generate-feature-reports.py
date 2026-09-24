#!/usr/bin/env python3
"""Generate per-feature Word deliverable reports in a single combined document.

Usage:
    python3 scripts/generate-feature-reports.py

Output:
    docs/TUES-Landing-Page-Feature-Reports.docx
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LOCAL_LIB = ROOT / ".venv-report-lib"
if LOCAL_LIB.is_dir():
    sys.path.insert(0, str(LOCAL_LIB))

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    import subprocess

    print("Installing python-docx to .venv-report-lib...", file=sys.stderr)
    LOCAL_LIB.mkdir(parents=True, exist_ok=True)
    subprocess.check_call(
        [sys.executable, "-m", "pip", "install", "python-docx", "-q", "--target", str(LOCAL_LIB)]
    )
    sys.path.insert(0, str(LOCAL_LIB))
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(ROOT / "scripts"))
from feature_reports_data import FEATURES, FeatureReport

OUTPUT = ROOT / "docs" / "TUES-Landing-Page-Feature-Reports.docx"
LEGACY_DIR = ROOT / "docs" / "feature-reports"
OXFORD_BLUE = RGBColor(0x00, 0x33, 0x66)
DOC_REF_PREFIX = "INV-EFF/TUES-PORTAL"


def set_cell_shading(cell, hex_color: str) -> None:
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml

    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_title(doc: Document, text: str, *, size: int = 16) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = OXFORD_BLUE


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.italic = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = OXFORD_BLUE


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        set_cell_shading(hdr[i], "E8EEF4")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def append_feature_sections(doc: Document, feature: FeatureReport) -> None:
    """Append one feature's content sections to an existing document."""
    doc_ref = f"{DOC_REF_PREFIX}/{feature.code}/2026/09/001"

    add_heading(doc, f"{feature.code} — {feature.title}", level=1)

    meta = [
        ("Document Reference", doc_ref),
        ("Parent Report", "INV-EFF/TUES-PORTAL/2026/09/001"),
        ("Delivered Effort", f"{feature.effort_md} ({feature.allocation_pct} of Module 1)"),
        ("Target Beneficiaries", feature.beneficiaries),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    add_heading(doc, "Business Purpose", level=2)
    doc.add_paragraph(feature.purpose)

    add_heading(doc, "Public Routes & URLs", level=2)
    doc.add_paragraph("The following routes are registered for this feature area:")
    for route in feature.routes:
        add_bullet(doc, route)

    if feature.admin_routes:
        add_heading(doc, "Admin CMS Routes", level=2)
        for route in feature.admin_routes:
            add_bullet(doc, route)

    add_heading(doc, "Source Files & Components", level=2)
    doc.add_paragraph("Key implementation files in the tues-landing-page repository:")
    for f in feature.source_files:
        add_bullet(doc, f)

    add_heading(doc, "Capabilities Delivered", level=2)
    for cap in feature.capabilities:
        add_bullet(doc, cap)

    if feature.known_gaps:
        add_heading(doc, "Known Gaps & Notes", level=2)
        for gap in feature.known_gaps:
            add_bullet(doc, gap)

    add_heading(doc, "Effort Summary", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Feature Code", feature.code],
            ["Feature Name", feature.title],
            ["Delivered Effort", feature.effort_md],
            ["Allocation (% of 35 MD)", feature.allocation_pct],
            ["Public Routes Listed", str(len(feature.routes))],
            ["Admin Routes", str(len(feature.admin_routes))],
            ["Source Files Referenced", str(len(feature.source_files))],
            ["Capabilities Delivered", str(len(feature.capabilities))],
        ],
        col_widths=[2.5, 4.0],
    )


def build_combined_document() -> Document:
    doc = Document()

    add_title(doc, "Per-Feature Deliverable Reports", size=18)
    add_subtitle(doc, "TUES University Public Web Portal & Admissions System")
    doc.add_paragraph()
    add_subtitle(doc, "Module 1 — PORTAL-1.1 through PORTAL-1.8")
    doc.add_paragraph()

    cover_meta = [
        ("Document Reference", "INV-EFF/TUES-PORTAL-FEATURES/2026/09/001"),
        ("Parent Report", "INV-EFF/TUES-PORTAL/2026/09/001"),
        ("Date of Issuance", "September 09, 2026"),
        ("Service Provider", "[Person]"),
        ("Client Organization", "[Place]"),
        ("Total Module 1 Effort", "35.0 Man-Days (280 Hours)"),
        ("Features Documented", str(len(FEATURES))),
    ]
    for label, value in cover_meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    doc.add_page_break()

    add_heading(doc, "Table of Contents")
    add_table(
        doc,
        ["Code", "Feature", "Effort", "Allocation"],
        [[f.code, f.title, f.effort_md, f.allocation_pct] for f in FEATURES]
        + [["TOTAL", "All PORTAL features", "35.0 MD", "100.0%"]],
        col_widths=[1.0, 3.2, 0.9, 0.9],
    )

    doc.add_page_break()

    for i, feature in enumerate(FEATURES):
        append_feature_sections(doc, feature)
        if i < len(FEATURES) - 1:
            doc.add_page_break()

    doc.add_page_break()
    add_heading(doc, "Formal Sign-Off")
    doc.add_paragraph(
        "IN WITNESS WHEREOF, the undersigned confirm delivery of the features "
        "detailed throughout this document:"
    )
    add_table(
        doc,
        ["", "SERVICE PROVIDER", "CLIENT"],
        [
            ["Signature", "___________________________", "___________________________"],
            ["Name", "[Person]", "[Person]"],
            ["Title", "Lead Developer", "Project Owner / CIO"],
            ["Date", "September 09, 2026", "___________________________"],
        ],
        col_widths=[1.2, 2.5, 2.5],
    )

    return doc


def remove_legacy_individual_files() -> None:
    if not LEGACY_DIR.is_dir():
        return
    for path in LEGACY_DIR.glob("PORTAL-*.docx"):
        path.unlink()
        print(f"Removed legacy: {path}")


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_combined_document()
    doc.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")

    remove_legacy_individual_files()

    index_lines = [
        "# TUES Landing Page — Per-Feature Deliverable Reports",
        "",
        "> Regenerate: `python3 scripts/generate-feature-reports.py`",
        "",
        "Combined document: [TUES-Landing-Page-Feature-Reports.docx](../TUES-Landing-Page-Feature-Reports.docx)",
        "",
        "Parent report: [TUES-Landing-Page-Deliverable-Report.docx](../TUES-Landing-Page-Deliverable-Report.docx)",
        "",
        "| Code | Feature | Effort |",
        "|------|---------|--------|",
    ]
    for f in FEATURES:
        index_lines.append(f"| {f.code} | {f.title} | {f.effort_md} |")

    LEGACY_DIR.mkdir(parents=True, exist_ok=True)
    index_path = LEGACY_DIR / "README.md"
    index_path.write_text("\n".join(index_lines) + "\n", encoding="utf-8")
    print(f"Updated: {index_path}")
    print(f"\nTotal: {len(FEATURES)} features in one document")


if __name__ == "__main__":
    main()

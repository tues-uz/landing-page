#!/usr/bin/env python3
"""Generate TUES Landing Page Deliverable Report as Word (.docx).

Usage:
    python3 scripts/generate-landing-page-report.py

Output:
    docs/TUES-Landing-Page-Deliverable-Report.docx
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LOCAL_LIB = ROOT / ".venv-report-lib"
if LOCAL_LIB.is_dir():
    sys.path.insert(0, str(LOCAL_LIB))

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    import subprocess

    print("Installing python-docx to .venv-report-lib...", file=sys.stderr)
    LOCAL_LIB.mkdir(parents=True, exist_ok=True)
    subprocess.check_call(
        [sys.executable, "-m", "pip", "install", "python-docx", "-q", "--target", str(LOCAL_LIB)]
    )
    sys.path.insert(0, str(LOCAL_LIB))
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

OUTPUT = ROOT / "docs" / "TUES-Landing-Page-Deliverable-Report.docx"

# Oxford Blue accent for headings
OXFORD_BLUE = RGBColor(0x00, 0x33, 0x66)


def set_cell_shading(cell, hex_color: str) -> None:
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml

    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = OXFORD_BLUE


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.italic = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = OXFORD_BLUE


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        set_cell_shading(hdr[i], "E8EEF4")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()

    # ── Cover ──
    add_title(doc, "Professional Services Invoice & Deliverable Effort Report")
    add_subtitle(doc, "TUES University Public Web Portal & Admissions System")
    doc.add_paragraph()
    add_subtitle(
        doc,
        "Executive Commercial Invoice & Business Deliverable Report:\n"
        "Portal Modernization, CMS Content Management, Admissions Intake & Multilingual Public Experience",
    )
    doc.add_paragraph()

    meta_items = [
        ("Document Reference", "INV-EFF/TUES-PORTAL/2026/09/001"),
        ("Date of Issuance", "September 09, 2026"),
        ("Project Title", "TUES University Public Web Portal & Admissions System"),
        ("Contract / PO Reference", "[File]"),
        ("Service Provider", "[Person]"),
        ("Client Organization", "[Place]"),
        (
            "Target Executive Audience",
            "University Steering Committee, Rectorate, Dean's Council, Finance Directorate, and IT Leadership",
        ),
        ("Total Delivered Effort", "35.0 Man-Days (280 Professional Engineering Hours)"),
        (
            "Scope Note",
            "This document covers Module 1 only — the tues-landing-page repository. "
            "EduHub LMS and the Academic Journal platform are separate deliverables.",
        ),
        (
            "Milestone Completion",
            "Core portal delivered; CMS operational; remaining items documented in Appendix A.",
        ),
    ]
    for label, value in meta_items:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    doc.add_page_break()

    # ── Section 1 ──
    add_heading(doc, "1. Executive Billing Summary & Commercial Schedule")
    add_heading(doc, "1.1 Executive Summary", level=2)
    doc.add_paragraph(
        "This document represents the formal Commercial Invoice and Business Deliverable Report submitted "
        "for the TUES University Public Web Portal. It provides institutional leadership with an executive-level "
        "account of the business objectives, operational workflows, and strategic value delivered for the "
        "public-facing digital presence and admissions intake system."
    )
    doc.add_paragraph(
        "The scope encompassed the complete conceptualization, branding, user experience design, software "
        "development, CMS integration, multilingual localization, quality verification on a managed staging "
        "environment, and packaging for handoff to the university's deployment infrastructure."
    )

    add_heading(doc, "1.2 Commercial Summary Recapitulation", level=2)
    add_table(
        doc,
        ["Indicator / Metric", "Value / Status", "Description / Notes"],
        [
            ["Total Delivered Effort", "35.0 Man-Days (MD)", "100.0% Module 1 Contract Delivery"],
            ["Productive Engineering Hours", "280.0 Hours", "Standard 8h / Man-Day equivalent"],
            ["Application Delivered", "1 Core Public Portal", "Public website + JWT-protected CMS admin"],
            ["Total Routes Configured", "155 Routes", "138 public, 15 admin, login, 404"],
            ["Total Page Components", "147 Page Files", "Plus 3 feature-routed pages"],
            ["Supported Languages", "4 (UZ / EN / RU / CN)", "48 locale JSON bundles + code fallbacks"],
            ["CMS Admin Modules", "9 Functional Modules", "Hero, news, events, programs, applications, newsletter, users"],
            ["Delivery Status", "Complete", "Core portal and CMS modules delivered"],
            ["Total Net Milestone Value", "Fixed-Price Milestone", "Net Amount as per Contract Schedule"],
        ],
        col_widths=[2.0, 1.8, 2.7],
    )

    # ── Section 2 ──
    add_heading(doc, "2. Strategic Business Value & Return on Investment (ROI)")
    doc.add_paragraph(
        "For the university's executive leadership and finance committee, this investment delivers "
        "tangible institutional benefits across six primary dimensions:"
    )
    roi_bullets = [
        "Revenue Acceleration: 24/7 global admissions intake via the digital application portal, with applications routed to the CMS admin inbox.",
        "Operational Efficiency: Self-service CMS for news, events, hero carousel, and study program content.",
        "Global Reach: Four-language public experience (Uzbek, English, Russian, Chinese) with localized study program titles and curriculum.",
        "Institutional Showcase: Interactive virtual campus tour (homepage Kuula embeds + Photo Sphere Viewer page at /virtual-tour).",
        "Academic Transparency: Dynamic bachelor and master's degree catalogs with downloadable program PDFs and curriculum detail pages.",
        "Statutory Accessibility: Floating accessibility widget with high-contrast, grayscale, inversion, font enlargement, and brightness/contrast controls.",
    ]
    for b in roi_bullets:
        add_bullet(doc, b)

    doc.add_page_break()

    # ── Section 3 ──
    add_heading(doc, "3. Module 1: TUES University Public Web Portal & Admissions System")
    p = doc.add_paragraph()
    p.add_run("Delivered Effort: ").bold = True
    p.add_run("35.0 Man-Days (280 Hours)\n")
    p.add_run("Target Beneficiaries: ").bold = True
    p.add_run("Prospective Students, Parents, International Partners, Alumni, Public Visitors, and University Communications Staff")

    add_heading(doc, "3.1 Business Purpose", level=2)
    doc.add_paragraph(
        "Transform the university's public digital identity into a comprehensive institutional showcase, "
        "driving international prospective student enrollments, centralizing campus communications, and "
        "providing staff with a modern content management workflow for news, events, and academic program information."
    )

    add_heading(doc, "3.2 Business Features & Operational Capabilities", level=2)
    portal_rows = [
        [
            "PORTAL-1.1",
            "Institutional Brand Identity & App Shell",
            "Oxford Blue/Gold visual identity, responsive header with mega-dropdown navigation, dynamic breadcrumbs, site-wide search, floating language switcher (UZ/EN/RU/CN), and footer.",
            "6.0 MD",
        ],
        [
            "PORTAL-1.2",
            "Interactive Virtual Campus Tour",
            "Homepage Kuula 360° embeds (3 collections) plus dedicated /virtual-tour page with Photo Sphere Viewer (linked nodes, gallery, custom tooltips).",
            "4.5 MD",
        ],
        [
            "PORTAL-1.3",
            "Digital Admissions 2025 Intake Portal",
            "Application form with study program selection, identity fields, captcha protection, and admin applications inbox with status workflow.",
            "5.5 MD",
        ],
        [
            "PORTAL-1.4",
            "Academic Programs & Curriculum Showcases",
            "Bachelor/masters hubs, CMS study programs (3 faculties, 37 programs), localized curriculum (509 courses), downloadable program PDFs.",
            "4.5 MD",
        ],
        [
            "PORTAL-1.5",
            "Executive Governance & Leadership Directory",
            "Organizational structure org-chart pages, dean/vice-rector/department leader profiles, university departments hub.",
            "3.5 MD",
        ],
        [
            "PORTAL-1.6",
            "Internationalization & Global Grants",
            "30 routes under /internationalization/* covering teacher exchange, global scholarships (MEXT, Matsumae, Humphrey), minority support.",
            "4.0 MD",
        ],
        [
            "PORTAL-1.7",
            "Digital Newsroom & Events Calendar",
            "Public news/events with date/category filters, CMS Tiptap editor, drag-and-drop ordering, event translations, newsletter management.",
            "4.0 MD",
        ],
        [
            "PORTAL-1.8",
            "Statutory Accessibility Toolbar (WCAG 2.1 AA)",
            "Floating widget: high contrast, grayscale, inversion, font size (+/− 4 steps), brightness/contrast; settings persisted in localStorage.",
            "3.0 MD",
        ],
        [
            "SUBTOTAL",
            "TUES PUBLIC WEB PORTAL",
            "155 Routes, 147 Pages, 4 Languages, 9 CMS Modules",
            "35.0 MD",
        ],
    ]
    add_table(
        doc,
        ["Sub-Module Code", "Feature / Capability Name", "Business Objective & User-Facing Functionality", "Effort"],
        portal_rows,
        col_widths=[1.0, 1.5, 3.3, 0.7],
    )

    add_heading(doc, "3.3 Technology Stack", level=2)
    add_table(
        doc,
        ["Layer", "Technologies"],
        [
            ["Frontend", "React 18, TypeScript 5, Vite 5"],
            ["Routing", "React Router v6 (155 routes)"],
            ["Styling", "Tailwind CSS 3, Radix UI / shadcn"],
            ["State / Data", "TanStack React Query 5"],
            ["Internationalization", "i18next, react-i18next (4 languages, 12 namespaces)"],
            ["CMS Rich Text", "Tiptap 3 (starter-kit, image, link, placeholder)"],
            ["Virtual Tour", "Photo Sphere Viewer 5 + Kuula embeds"],
            ["Animation", "GSAP 3 with ScrollTrigger"],
            ["Backend Integration", "REST API via VITE_API_BASE_URL, JWT auth for admin"],
            ["Accessibility", "Custom widget + open-accessibility library"],
        ],
        col_widths=[1.5, 5.0],
    )

    doc.add_page_break()

    # ── Section 4 ──
    add_heading(doc, "4. CMS Admin Console (Content Management Subsystem)")
    doc.add_paragraph(
        "The portal includes a JWT-authenticated admin console (/admin/*) enabling non-technical staff "
        "to manage public content without developer intervention."
    )
    add_table(
        doc,
        ["Admin Module", "Route", "Capabilities"],
        [
            ["Dashboard", "/admin", "Overview, quick actions, recent activity"],
            ["Hero Management", "/admin/hero", "Carousel slides + background video/image, UZ/EN/RU translations"],
            ["News Board", "/admin/news", "Drag-and-drop highlight ordering, article list"],
            ["News Articles", "/admin/news/articles", "Tiptap editor, cover image, UZ/EN/RU tabs, category/display type"],
            ["Events", "/admin/events", "Create/edit events, date/time/location/description, image upload, translations"],
            ["Study Programs", "/admin/study-programs", "Faculty manager, program CRUD, curriculum editing, translations"],
            ["Bachelor Programs", "/admin/bachelor-programs", "Full-time/correspondence program content editing"],
            ["Applications", "/admin/applications", "Admissions inbox, status updates, applicant detail view"],
            ["Newsletter", "/admin/newsletter", "Subscriber list management"],
            ["User Provisioning", "/admin/users", "Admin account creation with role-based permissions"],
        ],
        col_widths=[1.3, 1.8, 3.4],
    )

    # ── Section 5 ──
    add_heading(doc, "5. Consolidated Effort Allocation & Role Accounting")
    add_table(
        doc,
        ["Solution Domain / Sub-Module", "Delivered Effort (MD)", "Allocation (%)"],
        [
            ["PORTAL-1.1 Institutional Brand Identity & App Shell", "6.0 MD", "17.1%"],
            ["PORTAL-1.2 Interactive Virtual Campus Tour", "4.5 MD", "12.9%"],
            ["PORTAL-1.3 Digital Admissions 2025 Intake Portal", "5.5 MD", "15.7%"],
            ["PORTAL-1.4 Academic Programs & Curriculum Showcases", "4.5 MD", "12.9%"],
            ["PORTAL-1.5 Executive Governance & Leadership Directory", "3.5 MD", "10.0%"],
            ["PORTAL-1.6 Internationalization & Global Grants", "4.0 MD", "11.4%"],
            ["PORTAL-1.7 Digital Newsroom & Events Calendar", "4.0 MD", "11.4%"],
            ["PORTAL-1.8 Statutory Accessibility Toolbar", "3.0 MD", "8.6%"],
            ["TOTAL DELIVERED PROJECT EFFORT", "35.0 MD", "100.0%"],
        ],
        col_widths=[3.5, 1.5, 1.5],
    )

    doc.add_page_break()

    # ── Section 6 ──
    add_heading(doc, "6. Formal Billing Sign-Off & Payment Remittance Details")
    add_heading(doc, "6.1 Commercial Billing Summary", level=2)
    add_bullet(doc, "Total Delivered Professional Effort: 35.0 Man-Days (280 Productive Engineering Hours)")
    add_bullet(doc, "Contract Milestone: Public Web Portal Delivery & CMS Integration")
    add_bullet(doc, "Payment Terms: Net 14 Calendar Days from receipt of invoice")

    add_heading(doc, "6.2 Corporate Bank Wire Remittance Details", level=2)
    doc.add_paragraph("Please remit funds to the service provider's designated corporate bank account:")
    bank_items = [
        ("Account Name", "[Person]"),
        ("Account Number", "[File]"),
        ("Beneficiary Bank", "[Place]"),
        ("Bank Branch", "[Place]"),
        ("SWIFT / BIC Code", "[File]"),
        ("Payment Reference", "Invoice INV-EFF/TUES-PORTAL/2026/09/001"),
    ]
    for label, value in bank_items:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    add_heading(doc, "6.3 Formal Sign-Off", level=2)
    doc.add_paragraph(
        "IN WITNESS WHEREOF, the undersigned authorized representatives hereby confirm and certify "
        "the accuracy and delivery of the professional engineering efforts detailed throughout this dossier:"
    )
    add_table(
        doc,
        ["", "SUBMITTED BY (SERVICE PROVIDER)", "RECEIVED & VERIFIED BY (CLIENT)"],
        [
            ["Company", "[Person]", "[Place]"],
            ["Signature", "___________________________", "___________________________"],
            ["Name", "[Person]", "[Person]"],
            ["Title", "Managing Director / Lead Partner", "Chief Information Officer / Vice Rector"],
            ["Date", "September 09, 2026", "September 09, 2026"],
        ],
        col_widths=[1.0, 2.5, 2.5],
    )

    doc.add_page_break()

    # ── Appendix A ──
    add_heading(doc, "Appendix A — Remaining Work & Known Gaps")
    doc.add_paragraph(
        "This appendix documents items identified during delivery that are outside the core Module 1 scope "
        "or pending client content/decisions."
    )

    add_heading(doc, "A.1 Navigation Links Without Pages (68 items)", level=2)
    doc.add_paragraph(
        "The main navbar (About · Research · Admissions · News) and 22 Explore-more links are fully wired. "
        "The following 68 menu labels still resolve to # (no destination page):"
    )
    add_table(
        doc,
        ["Location", "Count", "Examples"],
        [
            ["Header social icons", "4", "Twitter, LinkedIn, Instagram, YouTube"],
            ["Mobile menu extras", "3", "Community, Colleges, Journal"],
            ["Explore more → University", "18", "License, Charter, Faculties, Open data, etc."],
            ["Explore more → Education", "4", "Course catalogue, Resources, Study plans, Syllabus"],
            ["Explore more → Science", "4", "Seminars, Scientific journals, Expected conferences, Academic council"],
            ["Explore more → Internationalization", "9", "International relations, grants hub, support center, etc."],
            ["Explore more → Student life", "1", "Contests"],
            ["Explore more → Admission 2025", "14", "Apply, FAQ, contract amounts, transfer info, etc."],
            ["Explore more → Information services", "6", "Latest news, video gallery, photo gallery, etc."],
            ["Explore more → Vacancies", "5", "Academic/admin/research positions, benefits"],
        ],
        col_widths=[2.0, 0.6, 3.9],
    )
    doc.add_paragraph("Source: CLIENT_DELIVERY_STATUS.md — wiring controlled by src/config/topNavHubData.ts")

    add_heading(doc, "A.2 Virtual Tour — Remaining Items", level=2)
    for item in [
        "Arrow SVG sizing refinement across breakpoints",
        "Controls layout review (caption / gallery / navbar ordering)",
        "Optional panorama asset optimization for low-end devices",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "A.3 External System Links", level=2)
    add_bullet(doc, "EduHub LMS (/eduhub) and Journal links in the header point to external systems not included in this repository.")
    add_bullet(doc, "Distance learning Explore-more link opens https://lms.tues.uz/ in a new tab.")

    add_heading(doc, "A.4 Content Enrichment (Not Blocking)", level=2)
    doc.add_paragraph(
        "Many routed pages exist with structural layout and i18n keys but may benefit from richer copy, "
        "images, or PDF attachments supplied by the university communications team. "
        "This is a content population task, not a development blocker."
    )

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
